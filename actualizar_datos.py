#!/usr/bin/env python3
"""
Genera bancos.js a partir del Excel "Directorio de Bancos FI".

Uso:
    pip install openpyxl
    python actualizar_datos.py ruta/al/Directorio_Bancos_FI.xlsx
    python actualizar_datos.py ruta/al/Directorio_Bancos_FI.xlsx --incluir-costos
    python actualizar_datos.py Directorio.xlsx --semen Cotizador_Donantes.xlsx   (agrega catálogo de semen Ovodonors)

Por defecto NO publica costos de proveedor (el sitio en GitHub Pages es público).
Usa --incluir-costos solo si el repositorio y el sitio son privados.
"""
import argparse, json, re, sys
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

OUT = Path(__file__).resolve().parent / "bancos.js"


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.]", "", str(v))
    return float(s) if s else None


def clean(v):
    return str(v).strip() if v not in (None, "") else ""


def link_of(row):
    return clean(row[1]).split(" ")[0].rstrip("/ ") if len(row) > 1 else ""


def fecha_of(row):
    return clean(row[9]) if len(row) > 9 else ""


def parse(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]

    # Localiza secciones por el título de la fila
    sections, cur = {}, None
    keys = {"OCP — FRESCO": "OCPF", "OCP — CONGELADO": "OCPC", "LAFER": "LAFER",
            "GENEVITY": "GENEVITY", "OVODONORS": "OVO"}
    for r in rows:
        first = clean(r[0]).upper()
        hit = next((k for t, k in keys.items() if t in first), None)
        if hit:
            cur = hit
            sections[cur] = []
            continue
        if first.startswith("⚠"):
            cur = None
        if cur:
            sections[cur].append(r)

    for k in keys.values():
        if k not in sections:
            sys.exit(f"No encontré la sección {k} en el Excel. Revisa los títulos.")

    fechas = {}

    def meta(k):
        r = next(r for r in sections[k] if clean(r[0]).lower().startswith("catálogo"))
        fechas[k] = fecha_of(r)
        return link_of(r)

    def body(k):
        # filas después del encabezado de tabla (la fila que sigue a "Catálogo / Link")
        rs = sections[k]
        i = next(i for i, r in enumerate(rs) if clean(r[0]).lower().startswith("catálogo"))
        return rs[i + 2:]

    def cat_order(c):
        m = re.search(r"\d+", c["c"]); return int(m.group()) if m else 99

    ocpf = {"link": meta("OCPF"), "cats": sorted(
        [{"c": clean(r[0]), "p": num(r[1])} for r in body("OCPF") if num(r[1])], key=cat_order)}

    ocpc = {"link": meta("OCPC"), "cats": sorted(
        [{"c": clean(r[0]), "cost": num(r[1]), "p": num(r[4])} for r in body("OCPC") if num(r[4])],
        key=cat_order)}

    ovo = {"link": meta("OVO"), "donors": [
        {"code": int(num(r[0])), "eggs": int(num(r[1])), "cost": num(r[2]),
         "pack": num(r[4]), "unit": num(r[5]), "av": clean(r[7])}
        for r in body("OVO") if num(r[0]) and num(r[4])]}

    lafer = {"link": meta("LAFER"), "origins": [
        {"o": clean(r[0]), "p": num(r[1]), "ship": num(r[2]) or 0, "t": clean(r[3])}
        for r in body("LAFER") if num(r[1])]}

    gen = {"link": meta("GENEVITY"), "origins": [
        {"o": clean(r[0]), "p": num(r[1]), "ship": 0,
         "t": " · ".join(x for x in [("Entrega " + clean(r[2])) if clean(r[2]) else "", clean(r[3])] if x)}
        for r in body("GENEVITY") if num(r[1])]}

    return ocpf, ocpc, ovo, lafer, gen, fechas


def parse_semen(path):
    """Lee el 'Cotizador de Donantes' (catálogo de semen OVODONORS)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Datos Donantes" not in wb.sheetnames or "Parámetros" not in wb.sheetnames:
        sys.exit("El archivo de semen debe tener las hojas 'Datos Donantes' y 'Parámetros'.")
    pr = wb["Parámetros"]
    envio = num(pr["B6"].value) or 0.0
    margen = num(pr["B9"].value) if pr["B9"].value is not None else 0.35
    iva = num(pr["B18"].value) if pr["B18"].value is not None else 0.16
    ws = wb["Datos Donantes"]
    head = [clean(c.value) for c in ws[3]]
    col = {h: i for i, h in enumerate(head)}

    def g(r, name):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    donors = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        pid = clean(g(r, "Perfil ID"))
        comp = num(g(r, "Compensación Económica (MXN)"))
        if not pid or not comp:
            continue
        precio = (comp + envio) / (1 - margen) * (1 + iva)
        donors.append({
            "id": pid.zfill(4), "edad": num(g(r, "Edad")), "etnia": clean(g(r, "Etnicidad (Principal)")),
            "etniaDet": clean(g(r, "Etnicidad (Detalle)")), "complexion": clean(g(r, "Complexión")),
            "piel": clean(g(r, "Color de Piel")), "ojos": clean(g(r, "Color de Ojos")),
            "cabello": clean(g(r, "Color de Cabello")), "tipoCabello": clean(g(r, "Tipo de Cabello")),
            "peso": num(g(r, "Peso (kg)")), "altura": num(g(r, "Altura (cm)")),
            "sangre": clean(g(r, "Tipo de Sangre")), "estudio": clean(g(r, "Área de Estudio")),
            "ocupacion": clean(g(r, "Ocupación")), "tatuajes": clean(g(r, "Tatuajes")),
            "documento": clean(g(r, "Documento Migratorio")), "pais": clean(g(r, "País de Nacimiento")),
            "categoria": clean(g(r, "Categoría de Donante")), "estatus": clean(g(r, "Estatus")) or "Disponible",
            "precio": round(precio), "cost": comp + envio,
        })
    return {"envio": envio, "margen": margen, "iva": iva, "donors": donors}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--semen", help="Excel 'Cotizador de Donantes' (catálogo de semen Ovodonors)")
    ap.add_argument("--incluir-costos", action="store_true")
    a = ap.parse_args()

    ocpf, ocpc, ovo, lafer, gen, f = parse(a.xlsx)

    semen = parse_semen(a.semen) if a.semen else None
    if semen is None and OUT.exists():
        # conserva el catálogo de semen ya publicado si no se pasa --semen
        prev = OUT.read_text(encoding="utf-8")
        try:
            semen = json.loads(prev[prev.index("{"):prev.rindex(";")]).get("SEMEN_OVO")
        except Exception:
            semen = None

    if not a.incluir_costos:
        if semen:
            for d in semen["donors"]:
                d["cost"] = None
        for c in ocpc["cats"]:
            c["cost"] = None
        for d in ovo["donors"]:
            d["cost"] = None

    data = {
        "DATA": {
            "OCPF": {"name": "OCP — Fresco", "short": "OCP Fresco", "type": "fresco",
                     "desc": "Paquete de óvulos por categoría", **ocpf},
            "OCPC": {"name": "OCP — Congelado", "short": "OCP Congelado", "type": "congelado",
                     "desc": "Precio por unidad", **ocpc},
            "OVO": {"name": "Ovodonors", "short": "Ovodonors", "type": "congelado",
                    "desc": "Por donante: lote o por óvulo", **ovo},
        },
        "SPERM": {"LAFER": {"name": "LAFER", **lafer}, "GENEVITY": {"name": "Genevity", **gen}},
        "SEMEN_OVO": semen,
        "META": {
            "fuente": "Directorio de Bancos FI",
            "actualizaciones": " · ".join(f"{n} {f[k]}" for n, k in
                                          [("OCP", "OCPF"), ("LAFER", "LAFER"), ("Ovodonors", "OVO")]
                                          if f.get(k) and f[k] != "—"),
            "generado": date.today().strftime("%d.%m.%Y"),
            "costos": a.incluir_costos,
        },
    }
    OUT.write_text("// Archivo generado por actualizar_datos.py. No editar a mano.\n"
                   "window.FI_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n",
                   encoding="utf-8")
    print(f"OK → {OUT}  ({len(ovo['donors'])} donantes de óvulos, "
          f"{len(semen['donors']) if semen else 0} donantes de semen, costos={'sí' if a.incluir_costos else 'no'})")


if __name__ == "__main__":
    main()
